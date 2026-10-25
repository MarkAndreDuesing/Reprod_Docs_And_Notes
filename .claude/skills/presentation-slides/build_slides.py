"""Generic slide builder (LMU house style) for summary and discussion slides.

Usage:
    python3 build_slides.py content.json out.pptx

The content file describes the deck (see content-template.json). Two modes:
  "new":    build a fresh 16:9 deck.
  "extend": open an existing deck ("base"), insert the new slides after the slide whose
            first text is "insert_after", and reuse that deck's layout and title shape
            (taken from the slide whose first text is "style_ref"). Other slides stay untouched.
"""
import copy, json, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def rgb(h): return RGBColor.from_string(h)

STYLE = {  # LMU-like house style; override per deck via "style"
    "font": "Arial", "accent": "00883A", "text": "1A1A1A", "muted": "5F6368",
    "panel": "F1F3F2", "panel_accent": "E3F1E8", "white": "FFFFFF",
    "header_pt": 40, "body_pt": 20, "small_pt": 16,
}
GROUPS = {  # default markers for categorised findings
    "A": {"symbol": "■", "color": "00883A", "name": "group A"},
    "B": {"symbol": "▲", "color": "5F6368", "name": "group B"},
    "C": {"symbol": "★", "color": "1A1A1A", "name": "group C"},
}


class Deck:
    def __init__(self, spec):
        self.s = {**STYLE, **spec.get("style", {})}
        self.groups = {**GROUPS, **spec.get("groups", {})}
        self.mode = spec.get("mode", "new")
        if self.mode == "extend":
            self.p = Presentation(spec["base"])
            ref = self.find(spec["style_ref"])
            self.layout = ref.slide_layout
            self.title_el = next(sh for sh in ref.shapes if sh.has_text_frame and sh.text_frame.text.strip())._element
            self.num_el = next((sh._element for sh in ref.shapes if sh.is_placeholder), None)
            self.anchor = self.find(spec["insert_after"])
        else:
            self.p = Presentation()
            self.p.slide_width, self.p.slide_height = Inches(13.333), Inches(7.5)
            self.layout = self.p.slide_layouts[6]
        self.new = []

    # ---------- helpers
    def find(self, text):
        for sl in self.p.slides:
            for sh in sl.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    if sh.text_frame.text.strip().startswith(text): return sl
                    break
        raise SystemExit(f"slide starting with {text!r} not found")

    def slide(self, header=None, dark=False):
        sl = self.p.slides.add_slide(self.layout)
        for ph in list(sl.placeholders): ph._element.getparent().remove(ph._element)
        if dark:
            bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = rgb(self.s["accent"])
        if header is not None:
            if self.mode == "extend":
                sl.shapes._spTree.append(copy.deepcopy(self.title_el))
                tf = sl.shapes[-1].text_frame
                tf.paragraphs[0].runs[0].text = header
                for r in tf.paragraphs[0].runs[1:]: r.text = ""
            else:
                self.text(sl, .7, .4, 11.93, .9, [[(header, {"b": True, "c": self.s["accent"], "size": self.s["header_pt"]})]])
        if self.mode == "extend" and self.num_el is not None:
            sl.shapes._spTree.append(copy.deepcopy(self.num_el))
        self.new.append(sl)
        return sl

    def text(self, sl, x, y, w, h, paras, size=None, anchor=None, align=None, space=6):
        """paras: list of paragraphs; a paragraph is a str or a list of (text, opts) runs."""
        box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.word_wrap = True
        if anchor: tf.vertical_anchor = anchor
        for i, runs in enumerate(paras):
            pa = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pa.space_after = Pt(space)
            pa.alignment = align or PP_ALIGN.LEFT
            if isinstance(runs, str): runs = [(runs, {})]
            for t, o in runs:
                r = pa.add_run(); r.text = t; f = r.font; f.name = self.s["font"]
                f.size = Pt(o.get("size", size or self.s["body_pt"])); f.bold = o.get("b", False)
                f.italic = o.get("i", False); f.color.rgb = rgb(o.get("c", self.s["text"]))
        return box

    def panel(self, sl, x, y, w, h, fill=None):
        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.08; sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill or self.s["panel"])
        sh.line.fill.background(); sh.shadow.inherit = False
        return sh

    def label(self, sl, t):
        self.text(sl, .7, 1.2, 9, .4, [[(t, {"c": self.s["muted"], "size": self.s["small_pt"], "b": True})]])

    def B(self, t, **o): return (t, {"b": True, **o})

    def labelled(self, item):
        """'text' or {'b': label, 't': text} -> runs"""
        if isinstance(item, str): return [(item, {})]
        return [self.B(item["b"] + " ", c=self.s["accent"]), (item["t"], {})]

    def sym(self, codes, size):
        return [(self.groups[c]["symbol"] + " ", {"c": self.groups[c]["color"], "size": size, "b": True}) for c in codes]

    def legend(self, sl, x=7.6, y=1.2, w=5.0):
        runs = []
        for k, g in self.groups.items():
            if g.get("hide"): continue
            runs += self.sym(k, 14) + [(g["name"] + "   ", {"c": self.s["muted"], "size": 14})]
        self.text(sl, x, y, w, .4, [runs], align=PP_ALIGN.RIGHT)

    def table(self, sl, header, rows, widths, y=1.7, size=17, row_h=.45, cell=None):
        t = sl.shapes.add_table(len(rows) + 1, len(header), Inches(.7), Inches(y), Inches(sum(widths)), Inches(row_h)).table
        for j, w in enumerate(widths): t.columns[j].width = Inches(w)
        for r in t.rows: r.height = Inches(row_h)
        for i, row in enumerate([header] + rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); c.fill.solid()
                c.fill.fore_color.rgb = rgb(self.s["accent"] if i == 0 else (self.s["panel"] if i % 2 else self.s["white"]))
                c.margin_left = c.margin_right = Inches(.1); c.margin_top = c.margin_bottom = Inches(.04)
                tf = c.text_frame; tf.text = ""
                paras = cell(i, j, val) if (cell and i > 0) else [[(str(val), {"b": i == 0, "c": self.s["white"] if i == 0 else self.s["text"]})]]
                for k, runs in enumerate(paras):
                    pa = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); pa.space_after = Pt(3)
                    for txt, o in runs:
                        r = pa.add_run(); r.text = txt; f = r.font; f.name = self.s["font"]
                        f.size = Pt(o.get("size", size)); f.bold = o.get("b", False); f.color.rgb = rgb(o.get("c", self.s["text"]))
        return t

    def notes(self, sl, d):
        if d.get("notes"): sl.notes_slide.notes_text_frame.text = d["notes"]

    # ---------- slide types
    def t_title(self, d):
        sl = self.slide(dark=True)
        self.text(sl, .7, 2.4, 11.9, 1.2, [[self.B(d["title"], c=self.s["white"], size=40)]])
        if d.get("subtitle"): self.text(sl, .7, 3.6, 11.9, 1, [[(d["subtitle"], {"c": self.s["white"], "size": 20})]])
        if d.get("small"): self.text(sl, .7, 4.6, 11.9, .6, [[(d["small"], {"c": self.s["white"], "size": 14})]])
        return sl

    def t_section(self, d):
        sl = self.slide(dark=True)
        self.text(sl, .7, 2.6, 11.9, 1, [[self.B(d["title"], c=self.s["white"], size=40)]])
        if d.get("subtitle"): self.text(sl, .7, 3.6, 9.5, 1.2, [[(d["subtitle"], {"c": self.s["white"], "size": 18})]])
        return sl

    def t_bullets(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        self.text(sl, .7, 1.7, 11.93, 5.2, [[("• ", {})] + self.labelled(b) for b in d["bullets"]], space=12)
        return sl

    def t_callout(self, d):
        sl = self.slide(d["title"])
        self.text(sl, .7, 1.75, 3.3, 4.8, [[self.B(c, c=self.s["accent"], size=24)] for c in d["callout"]], space=18)
        self.text(sl, 4.34, 1.6, 8.3, 5.2, [self.labelled(b) for b in d["bullets"]], size=19, space=12)
        return sl

    def t_rows(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        n = len(d["rows"]); top = 1.6; h = min(1.15, (5.6 - .15 * (n - 1)) / n)
        for i, (k, v) in enumerate(d["rows"]):
            y = top + i * (h + .15)
            self.panel(sl, .7, y, 11.93, h)
            self.text(sl, .9, y + .05, 3.0, h - .1, [[self.B(k, c=self.s["accent"], size=20)]], anchor=MSO_ANCHOR.MIDDLE)
            self.text(sl, 3.9, y + .05, 8.6, h - .1, [[(v, {"size": 18})]], anchor=MSO_ANCHOR.MIDDLE)
        return sl

    def t_table(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        widths = d.get("widths") or [11.93 / len(d["header"])] * len(d["header"])
        self.table(sl, d["header"], d["rows"], widths, size=d.get("size", 17), row_h=.5)
        if d.get("foot"): self.text(sl, .7, 6.4, 11.93, .6, [[(d["foot"], {"c": self.s["muted"], "size": 15})]])
        return sl

    def t_tiles(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        y = 1.7
        if d.get("statement"):
            self.panel(sl, .7, y, 11.93, 1.5, self.s["panel_accent"])
            self.text(sl, .9, y + .05, 11.5, 1.4, [self.labelled(x) for x in d["statement"]], size=19, anchor=MSO_ANCHOR.MIDDLE)
            y += 1.8
        if d.get("tiles_label"):
            self.text(sl, .7, y, 11.93, .4, [[self.B(d["tiles_label"], c=self.s["muted"], size=16)]]); y += .5
        n = len(d["tiles"]); w = (11.93 - .2 * (n - 1)) / n
        for i, (k, v) in enumerate(d["tiles"]):
            x = .7 + i * (w + .2)
            self.panel(sl, x, y, w, 2.4, self.s["panel_accent"] if not d.get("statement") else None)
            self.text(sl, x + .15, y + .15, w - .3, 2.1, [[self.B(k, c=self.s["accent"], size=20)], [(v, {"size": 17})]], space=10)
        if d.get("foot"): self.text(sl, .7, 6.4, 11.93, .5, [[(d["foot"], {"size": 18})]])
        return sl

    def t_panels(self, d):
        sl = self.slide(d["title"])
        y = 1.4
        if d.get("lead"):
            self.text(sl, .7, y, 11.93, .6, [self.labelled(d["lead"])], size=19); y += .75
        n = len(d["panels"]); w = (11.93 - .3 * (n - 1)) / n
        for i, pnl in enumerate(d["panels"]):
            x = .7 + i * (w + .3)
            self.panel(sl, x, y, w, 6.15 - y, self.s["panel_accent"])
            self.text(sl, x + .2, y + .1, w - .4, 6.0 - y,
                      [[self.B(pnl["title"], c=self.s["accent"], size=21)]] + [[("• " + b, {"size": 18})] for b in pnl["bullets"]], space=10)
        if d.get("foot"): self.text(sl, .7, 6.3, 11.93, .5, [[self.B("→ ", c=self.s["accent"]), (d["foot"], {})]], size=18)
        return sl

    def t_grid(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        cols = d.get("cols", 4); w = (11.93 - .2 * (cols - 1)) / cols
        for i, it in enumerate(d["items"]):
            x = .7 + (i % cols) * (w + .2); y = 1.75 + (i // cols) * 1.5
            self.panel(sl, x, y, w, 1.3)
            self.text(sl, x + .12, y + .08, w - .25, 1.15,
                      [[self.B(it["id"] + " ", c=self.s["accent"], size=18), (it["name"], {"size": 17})], self.sym(it.get("groups", ""), 18)], space=4)
        self.legend(sl, .7, 6.35, 11.93); sl.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        return sl

    def t_category_table(self, d):
        sl = self.slide(d["title"])
        if d.get("label"): self.label(sl, d["label"])
        self.legend(sl)
        rows = [[it, it] for it in d["items"]]
        def cell(i, j, it):
            if j == 0:
                return [[self.B(it["id"] + " ", c=self.s["accent"]), self.B(it["name"])], self.sym(it.get("groups", ""), 16)]
            return [self.sym(g, 16) + [(t, {"size": 16})] for g, t in it["points"]]
        self.table(sl, d.get("header", ["Category", "Points raised"]), rows, d.get("widths", [2.7, 9.23]), cell=cell)
        if d.get("foot"): self.text(sl, .7, 6.55, 11.93, .45, [[(d["foot"], {"c": self.s["muted"], "size": 14})]])
        return sl

    def t_question(self, d):
        sl = self.slide(d.get("header", "Discussion"))
        self.text(sl, .7, 1.45, 11.93, .7, [[self.B(d["title"], c=self.s["accent"], size=26)]], space=4)
        self.text(sl, .7, 2.1, 11.93, .8, [[(d["context"], {"c": self.s["muted"], "size": 20})]], space=4)
        self.text(sl, .7, 2.95, 11.93, 1.8, [[self.B(d["question"], size=28)]], space=4)
        ups = [(k, d[k]) for k in ("example", "perspective") if d.get(k)]
        labels = {"example": "Example", "perspective": d.get("perspective_label", "You")}
        if ups:
            h = .55 + .65 * len(ups)
            self.panel(sl, .7, 7.0 - h, 11.93, h - .2)
            self.text(sl, .9, 7.1 - h, 11.5, h - .4,
                      [[self.B(f"↳ {labels[k]}: ", c=self.s["accent"]), (v, {})] for k, v in ups], size=20, space=8)
        parts = []
        if d.get("answers"): parts.append("EXAMPLE ANSWERS\n" + "\n".join("- " + a for a in d["answers"]))
        if d.get("background"): parts.append("BACKGROUND (not on slide)\n" + "\n".join("- " + a for a in d["background"]))
        if parts: sl.notes_slide.notes_text_frame.text = "\n\n".join(parts)
        return sl

    def t_references(self, d):
        sl = self.slide(d.get("title", "References"))
        self.text(sl, .7, 1.6, 11.93, 5, [[self.B(f"[{i}] "), (r, {})] for i, r in enumerate(d["refs"], 1)], size=18, space=14)
        return sl

    # ---------- run
    def build(self, slides):
        for d in slides:
            sl = getattr(self, "t_" + d["type"])(d)
            if d["type"] != "question": self.notes(sl, d)
            if d.get("hidden"): sl._element.set("show", "0")
        if self.mode == "extend":
            lst = self.p.slides._sldIdLst
            eid = lambda s: next(e for e in lst if self.p.part.rels[e.rId]._target is s.part)
            prev = eid(self.anchor)
            for s in self.new:
                e = eid(s); lst.remove(e); prev.addnext(e); prev = e


if __name__ == "__main__":
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    deck = Deck(spec); deck.build(spec["slides"]); deck.p.save(sys.argv[2])
    print(f"{len(deck.new)} slides written to {sys.argv[2]} ({len(deck.p.slides)} in deck)")
