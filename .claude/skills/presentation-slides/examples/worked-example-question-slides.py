"""Paper 2 (Timperley) discussion questions: single source for the notes file and the slides.

Usage: python3 discussion_questions_timperley.py <in.pptx> <out.pptx> <out.md>
Inserts the question slides, sorted by challenge, after the slide titled
"Paper 2 · Discussion (Kevin)". The existing idea-dump slides follow and are hidden.
"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Each item: (title, paper context, question, example, you, example answers, my notes)
Q = [
 ("C1 Not worth it: Fear of exposing mistakes",
  "Creators fear shared code exposes mistakes and harms their reputation, though finding mistakes is what review is for.",
  "What would make you more or less willing to share your own code?",
  "a bug found later only leads to a correction note, not a retraction",
  "Would you share more readily if messy, as-is code were the norm?",
  ["More willing: proportionate consequences for honest errors; messy code accepted.",
   "Less willing: as long as any visible mistake can cost a paper or reputation."],
  ["Chang 2006: inherited script flipped two columns, 5 retractions, treated as honest error.",
   "Reinhart & Rogoff 2013: spreadsheet error found after the data was shared.",
   "Hermann et al.: some reviewers prefer availability over quality."]),
 ("C1 Not worth it: Holding back artifacts",
  "Some creators hold back artifacts to protect future publications.",
  "Does this slow down research, and how far should researchers accept being undercut for scientific progress?",
  "a tool kept private until the PhD student's next paper",
  "Would you build on a paper whose tool you can't get?",
  ["Protect: a PhD has fixed time; being overtaken with your own tool is a real risk.",
   "Share: the paper already establishes priority; whole research areas died without artifacts."],
  ["Embargoes in data-heavy fields (telescope data 12-18 months); no measured scooping rate."]),
 ("C2 Portability: Running untrusted artifacts",
  "Running untrusted artifacts is a security risk for reviewers.",
  "Should reviewers be allowed to refuse artifacts that need root access or invasive libraries?",
  "an installer that changes system libraries on the reviewer's laptop",
  "If refusing were allowed, which research could no longer be evaluated?",
  ["Yes: safety first; authors must package safely (VM, container).",
   "No: system-level research (OS, drivers, security tools) would become unevaluable."],
  ["POPL recommends VM images partly to protect reviewers' machines."]),
 ("C3 Maintenance: Decay and maintenance",
  "Artifacts decay as dependencies age and student maintainers leave.",
  "Is software decay unavoidable, and what practical ways are there to limit it?",
  "a 2015 thesis tool needs Python 2 and a library that no longer exists: fix it, freeze it in an image, or let it go?",
  "Who should pay for keeping your thesis tool running after you leave?",
  ["Unavoidable: you can slow decay, not stop it.",
   "Limitable: a frozen image runs as long as it can be executed.",
   "Cost: group or university vs. nobody (archive and mark unmaintained)."],
  ["Paper's own suggestions (teammate): maintenance plans, long-term archives.",
   "DFG: archive code and data 10 years; archived is not runnable.",
   "Research software engineers; short-cycle maintenance funding.",
   "Week 1: Docker may not exist in 40 years; source code most stable."]),
 ("C4 Tacit knowledge: One-click reproduction scripts",
  "Scripts capture steps that would otherwise stay in the developer's head. Reviewers warn they can be a black box that re-runs the authors' mistakes.",
  "Is a one-click reproduction script good or bad for scientific scrutiny?",
  "the script reproduces every table, but nobody reads the code behind it",
  "As a reviewer, what besides a script would help you understand why the results come out?",
  ["Good: anyone can re-run it.", "Bad: checks the numbers, not the method."],
  ["Python file-order case (2019): scripts correct on one OS, wrong on another."]),
 ("C4 Tacit knowledge: Capturing tacit knowledge",
  "Much know-how about a prototype stays in the developer's head.",
  "What practical ways are there to capture this knowledge so that others can use the artifact?",
  "the order the scripts must run in, hard-coded paths, parameter values that were tried and dropped",
  "Could someone reproduce your bachelor thesis results from your repository alone?",
  ["Mostly capturable with checklists and a run-all script.",
   "Not fully: some know-how only transfers by talking to the author."],
  []),
 ("C4 Tacit knowledge: Reviewer effort",
  "Reviewers report missing instructions and needing the paper to fill gaps; deep dives can take days.",
  "How much effort, time and money should reviewers invest before giving up on an artifact?",
  "after two hours it fails on a missing dependency: fix it, ask the authors, or reject?",
  "As an author, how much of a reviewer's goodwill would you rely on?",
  ["Little: it should work for a stranger.", "More: reviewers are often the only ones who ever try."],
  ["Typical budget 8-20 h per artifact (SOSP 2021, APLAS 2022)."]),
 ("C7 Hosting: Dead links and old badges",
  "Many artifacts become inaccessible, older ones more often.",
  "Should a badge expire if the artifact link dies after publication?",
  "conferences re-check links every few years and mark dead ones",
  "Would \"awarded 2016, not re-checked\" change how you treat a paper?",
  ["Yes: a badge should describe the current state.", "No: it certifies what was checked at review time."],
  ["Who organises and pays for re-checks?", "Klein et al. 2014: 13-22 % of web references suffer link rot."]),
 ("C7 Hosting: Very large artifacts",
  "Large datasets don't fit common hosting or a reviewer's disk.",
  "What should researchers do with artifacts too big for Zenodo or a reviewer's laptop?",
  "several terabytes of mined repository data",
  "What would you publish if your thesis data were that large?",
  ["A representative subset plus scripts to regenerate the rest.", "Institutional repositories; share the generator, not the data."],
  ["Zenodo: 50 GB per record, up to 200 GB on request.",
   "Overlaps with the paper's proof-of-concept recommendation (teammate)."]),
 ("C9 External constraints: Unshareable industry data",
  "Industry partners often forbid sharing data.",
  "Should reviewers criticise a paper for not sharing data its industry partner won't release?",
  "a study of a company's bug reports, with the data under NDA",
  "If you built on a study whose data is under NDA, would you trust it less than one whose artifact simply disappeared?",
  ["Yes: readers cannot check the claims either way.", "No: it punishes industry collaboration."],
  ["ICSE: sharing not required for acceptance, only a statement of why not."]),
 ("C9 External constraints: Formal exceptions",
  "ACM separates \"evaluated\" from \"publicly available\".",
  "Should artifact evaluation have formal exceptions for artifacts like industry data?",
  "reviewers get temporary access under NDA, but readers never see the artifact",
  "Would a badge saying \"evaluated privately\" convince you as a reader?",
  ["Yes: someone independent has checked it.", "No: a check nobody can repeat adds little."],
  ["ACM: Evaluated badges with reviewer-only access; Available needs a public archive. Proxies allowed for proprietary data."]),
 ("C9 External constraints: Open by default",
  "IP, privacy and commercial tools limit what can be shared.",
  "How open should research software be, and where are the legitimate limits?",
  "a company co-authored the tool and wants to sell it later",
  "Which part of your own thesis work would you not want to make public, and why?",
  ["Fully open: publicly funded research belongs to the public.",
   "Limits: personal data, partner contracts, commercial plans."],
  ["Horizon Europe: open data mandatory, exceptions justified in a data management plan."]),
 ("C9 External constraints: Licensed software",
  "Some artifacts need paid licences reviewers don't have.",
  "Should research tools that depend on paid commercial software be accepted as artifacts at all?",
  "your artifact needs MATLAB, and the reviewer has no licence: whose problem is that?",
  "As a reviewer, would you pay for a licence yourself to evaluate someone else's artifact?",
  ["Accept: much real-world SE runs on commercial platforms.", "Restrict: if others cannot run it, it misses its purpose."],
  ["ACM policy is silent on reviewer licences; each AEC decides. CAV/FORMATS/QEST: case by case, e.g. MATLAB Online."]),
 ("C10 Lack of reviewer incentives: Students as reviewers",
  "Artifact review is poorly rewarded and often left to students.",
  "Is leaving artifact evaluation to students a good training opportunity, or does it undermine how rigorous and seriously taken the badges are?",
  "a first-year PhD student judges an experienced group's artifact",
  "Would you trust a badge more if you knew a senior researcher rather than a student had awarded it?",
  ["Meaningful: checking that it runs needs no seniority.", "Weak: students may not judge methods or dare to reject."],
  ["Students chosen deliberately from ECOOP/OOPSLA 2013; most AEC members serve once (Hermann et al.)."]),
 ("C11 Technical obstacles: Fixed time limits",
  "Some artifacts need more than the roughly one day a reviewer has.",
  "Is it fair to fail an artifact after a fixed time limit, e.g. a 30-minute setup?",
  "a fuzzing experiment that runs for three days",
  "Could your own thesis project be set up in 30 minutes?",
  ["Fair: future users won't invest more either.", "Unfair: penalises inherently heavy research."],
  ["ICSE 2026: setup over 30 min makes acceptance unlikely; MobiSys allows a reduced representative scope."]),
 ("C11 Technical obstacles: Hardware reviewers don't have",
  "Reviewers often lack the hardware to run an artifact.",
  "Should conferences provide compute and storage for reviewers, e.g. remote servers with containers?",
  "an experiment needs eight GPUs for a week, and the reviewer has a laptop",
  "Would you trust a badge from a reviewer who only ran a scaled-down version?",
  ["Yes: otherwise resource-rich labs are favoured.", "No: costly; authors should provide access or a smaller run."],
  ["Paper recommends compute funding (teammate's section): keep it short and hand over.",
   "SOSP 2019 Azure credits proved hard to use; IISWC has authors provide remote access."]),
 ("C12 Limited communication: One-shot or iterative review",
  "Reviewers value responsive authors; single-round reviews frustrate both sides.",
  "Should artifact evaluation be a one-shot review or a dialogue with the authors?",
  "the artifact fails on a typo in the install instructions: reject, or ask the authors?",
  "Later users won't have the authors' help: should reviewers?",
  ["Dialogue: fixes help every later user.", "One-shot: a badge should mean it works without the authors."],
  ["Most venues have a 'kick the tires' phase; ICSE 2026: answers within 72 h."]),
 ("C12 Limited communication: Support after publication",
  "Users value authors who help fix problems.",
  "Is that support sustainable, given it can't last forever?",
  "the student who built the tool graduated two years ago, and a user's bug report goes unanswered",
  "Would you answer emails about your thesis code in three years?",
  ["Yes, for a stated period.", "No: archive it and say it is unsupported."],
  ["Overlaps with C3."]),
 ("RQ1: Rewards",
  "Respondents want sharing encouraged and rewarded.",
  "Which reward would actually make researchers invest more in artifacts?",
  "a best-artifact award at the conference vs. your artifacts counting in your next job application",
  "What would make you spend an extra week on your thesis artifact?",
  ["Career credit.", "Direct benefit: users, collaborations."],
  ["JOSS; possible overlap with the teammate's recommendations."]),
]
SPARES = []

ORDER = [x[0] for x in Q]

def ordered():
    items = {x[0]: x for x in Q + SPARES}
    assert set(items) == set(ORDER), set(items) ^ set(ORDER)
    return [items[t] for t in ORDER]

def markdown():
    out = ["# Timperley et al. (2021): Discussion Questions (Challenges, §4.2)", "",
           "Generated from `Vorl_2/slides/discussion_questions_timperley.py` (single source for slides and notes).", "",
           "Core questions come from my selected (green) questions, lightly reworded and anchored in the paper. "
           "Follow-ups go on the slide as sub-points. Example answers and my notes go in the speaker notes only; "
           "the notes draw on `timperley-research.md` and are not part of the questions.", ""]
    def block(items, hdr):
        out.append(f"## {hdr}"); out.append("")
        for i, (t, ctx, q, ex, me, a, n) in enumerate(items, 1):
            out.extend([f"### {i}. {t}", "", f"**Context:** {ctx}", "", f"**Question:** {q}", "",
                        "**Follow-ups:**", f"- Example: {ex}", f"- You: {me}", ""])
            out.extend(["**Example answers:**"] + [f"- {x}" for x in a] + [""])
            out.extend(["**My notes:**"] + [f"- {x}" for x in n] + [""])
    block(ordered(), "Questions, sorted by challenge (RQ1 last)")
    return "\n".join(out)

def build(src, dst):
    G = RGBColor(0x00, 0x88, 0x3A); K = RGBColor(0x1A, 0x1A, 0x1A); M = RGBColor(0x5F, 0x63, 0x68)
    p = Presentation(src)
    slides = list(p.slides)
    def ttl(s): return next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
    anchor = next(s for s in slides if ttl(s).startswith("Paper 2 · Discussion (Kevin)"))
    dumps = [s for s in slides[slides.index(anchor) + 1:slides.index(anchor) + 4] if ttl(s) == "Paper 2 · Discussion"]
    ref = next(s for s in slides if ttl(s) == "General Summary" and slides.index(s) > 10)
    title_el = next(sh for sh in ref.shapes if sh.has_text_frame and sh.text_frame.text == "General Summary")._element
    num_el = next(sh for sh in ref.shapes if sh.is_placeholder)._element
    def add(item, hidden):
        t, ctx, q, ex, me, a, n = item
        f = ["Example: " + ex, "You: " + me]
        s = p.slides.add_slide(ref.slide_layout)
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        s.shapes._spTree.append(copy.deepcopy(title_el)); s.shapes[-1].text_frame.paragraphs[0].runs[0].text = "Paper 2 · Discussion"
        s.shapes._spTree.append(copy.deepcopy(num_el))
        def para(tf, first, text, size, bold=False, color=K, space=8, prefix=None):
            pa = tf.paragraphs[0] if first else tf.add_paragraph(); pa.space_after = Pt(space)
            if prefix:
                r = pa.add_run(); r.text = prefix; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = G; r.font.name = "Arial"
            r = pa.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Arial"
        box = s.shapes.add_textbox(Inches(.7), Inches(1.45), Inches(11.93), Inches(1.0)); tf = box.text_frame; tf.word_wrap = True
        para(tf, True, t, 26, True, G, 4)
        box = s.shapes.add_textbox(Inches(.7), Inches(2.1), Inches(11.93), Inches(.8)); tf = box.text_frame; tf.word_wrap = True
        para(tf, True, ctx, 20, False, M, 4)
        box = s.shapes.add_textbox(Inches(.7), Inches(2.95), Inches(11.93), Inches(1.8)); tf = box.text_frame; tf.word_wrap = True
        para(tf, True, q, 28, True, K, 4)
        if f:
            h = 0.55 + 0.65 * len(f)
            pn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.7), Inches(7.0 - h), Inches(11.93), Inches(h - .2))
            pn.adjustments[0] = 0.08; pn.fill.solid(); pn.fill.fore_color.rgb = RGBColor(0xF1, 0xF3, 0xF2); pn.line.fill.background(); pn.shadow.inherit = False
            box = s.shapes.add_textbox(Inches(.9), Inches(7.0 - h + .1), Inches(11.5), Inches(h - .4)); tf = box.text_frame; tf.word_wrap = True
            for i, x in enumerate(f):
                lab, rest = x.split(": ", 1)
                para(tf, i == 0, rest, 20, False, K, 8, "↳ " + lab + ": ")
        s.notes_slide.notes_text_frame.text = ("EXAMPLE ANSWERS\n" + "\n".join("- " + x for x in a) +
                                               "\n\nMY NOTES (background, not on slide)\n" + ("\n".join("- " + x for x in n) or "-"))
        if hidden: s._element.set("show", "0")
        return s
    new = [add(x, False) for x in ordered()]
    for d in dumps: d._element.set("show", "0")
    lst = p.slides._sldIdLst
    def eid(sl): return next(e for e in lst if p.part.rels[e.rId]._target is sl.part)
    prev = eid(anchor)
    for s in new:
        e = eid(s); lst.remove(e); prev.addnext(e); prev = e
    p.save(dst)
    print("slides", len(p.slides), "dumps hidden", len(dumps))

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
    open(sys.argv[3], "w").write(markdown() + "\n")
