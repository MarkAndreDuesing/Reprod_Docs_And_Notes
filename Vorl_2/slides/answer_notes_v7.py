"""Adds answer plans to the speaker notes of the visible Paper 2 question slides in deck v7,
without changing any slide content. Existing notes are kept below the plan."""
import sys, importlib.util
spec = importlib.util.spec_from_file_location("v6", __file__.replace("answer_notes_v7.py", "answer_notes_v6.py"))
src_txt = open(spec.origin).read().split("src, dst = sys.argv[1], sys.argv[2]")[0]
ns = {}; exec(src_txt, ns); N = ns["N"]
for d in N.values(): d.pop("question", None)
N.pop("C7 Hosting / C11 Technical obstacles: Unreviewable artifacts")
N.pop("C12 Limited communication: Purpose of artifact review")
N["C7 Hosting: Very large artifacts"] = dict(
 about="Some datasets or artifacts are too big for common archives (Zenodo holds 50 GB per record) and for a reviewer's laptop. How can they still be shared and checked?",
 answers=[
  "SUBSET - share a small, representative sample plus the scripts that produced the full data; reviewers check the pipeline on the sample.",
  "SHARE THE RECIPE, NOT THE DATA - if the data comes from public sources (e.g. mined repositories), publish the exact scripts, queries and dates so others can regenerate it.",
  "OTHER HOSTING - university or national research data repositories, or splitting into well-documented parts; costs may need to be in the project budget.",
  "REVIEWER ACCESS - authors keep the data on their own server and give reviewers temporary remote access; readers later get only the subset.",
  "ACCEPT LIMITS - some artifacts simply can't be fully reviewed; the badge or paper should then say what was checked.",
  "DISCUSSION ANGLE - does size make an artifact effectively unreviewable, and should that affect how much we trust the paper?"],
 ladder=[
  "Several terabytes of mined repository data: what would you publish?",
  "What would you publish if your thesis data were that large?",
  "If reviewers can only check a sample, should the badge say so?"])
N["C11 Technical obstacles: Hardware reviewers don't have"] = dict(
 about="Some experiments need hardware reviewers don't have (many GPUs, clusters, special devices). Who should make the resources available, and what does a review still mean otherwise?",
 answers=[
  "CONFERENCE PROVIDES - conferences or ACM/IEEE fund shared machines or cloud credits, so reviewing doesn't depend on where a reviewer works (this is also one of the paper's recommendations; keep short, Minh Thi covers it).",
  "AUTHORS PROVIDE - authors give reviewers remote access to their own hardware or pay for cloud time; this is what most venues ask today.",
  "SCALED-DOWN RUN - authors provide a smaller version that runs on a normal machine and shows the same trend; weaker but practical.",
  "NO BADGE - if nobody can run it, nobody can certify it; say 'not evaluated' honestly.",
  "FAIRNESS - if authors must pay, well-funded groups get badges more easily; if conferences pay, someone has to fund it.",
  "TRUST - is a badge after a scaled-down run still meaningful to readers?"],
 ladder=[
  "An experiment needs eight GPUs for a week, and the reviewer has a laptop: what now?",
  "Would you trust a badge from a reviewer who only ran a scaled-down version?",
  "As an author, would you pay for reviewers' cloud time to get a badge?"])
N["C12 Limited communication: One-shot or iterative review"] = dict(
 about="Some artifact committees review in one round, others let reviewers and authors talk and fix problems. Which is better, and what does that say about what a badge should mean?",
 answers=[
  "ITERATIVE - small glitches (a typo in install instructions) are inevitable; fixing them during review improves the artifact for everyone later.",
  "ONE-SHOT - later users won't have the authors on call; the review should test whether it works on its own, like it will for them.",
  "HYBRID - a short 'does it start' phase with author help, then a final test without help; fixes must go into the artifact itself.",
  "WHAT IS REVIEW FOR - confirming the paper's results once (peer review of evidence) vs. certifying a tool others can build on (the other paper: reviewers want reproducible results, users want reusable tools).",
  "SUSTAINABILITY - authors can't support a tool forever; students graduate, emails go unanswered; a badge can only describe the state at review time."],
 ladder=[
  "The artifact fails on a typo in the install instructions: reject, or ask the authors?",
  "Later users won't have the authors' help: should reviewers?",
  "If you started a thesis on a badged tool, what would you expect the badge to guarantee?",
  "Would you still answer emails about your thesis code in three years?"])

src, dst = sys.argv[1], sys.argv[2]
from pptx import Presentation
p = Presentation(src); done = set()
for s in p.slides:
    if s._element.get("show") == "0": continue
    subs = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text in N]
    if not subs: continue
    key = subs[0].text_frame.text; d = N[key]; done.add(key)
    old = s.notes_slide.notes_text_frame.text.strip()
    s.notes_slide.notes_text_frame.text = (
        "WHAT THIS IS ABOUT\n" + d["about"] +
        "\n\nPOSSIBLE ANSWERS\n" + "\n".join("- " + a for a in d["answers"]) +
        "\n\nIF THE ROOM IS QUIET (ask in this order)\n" + "\n".join(f"{i}. {x}" for i, x in enumerate(d["ladder"], 1)) +
        ("\n\nEARLIER NOTES (kept)\n" + old if old else ""))
print("missing:", set(N) - done)
p.save(dst)
