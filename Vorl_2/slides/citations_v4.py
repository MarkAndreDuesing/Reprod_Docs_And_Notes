"""Adds [n] citations for all claims not from Timperley et al. (2021) on visible slides 11-42
(slide text and speaker notes) and rewrites the References slide."""
import sys
from pptx import Presentation
from pptx.util import Pt

REFS = [
 "ACM: Artifact Review and Badging, Version 1.1 (2020). acm.org/publications/policies/artifact-review-and-badging-current",
 "Winter & Beyer: Introductory Lecture, Reproducibility of Scientific Results in SE Research, LMU (2026)",
 "ICSE 2026 Artifact Evaluation, call for submissions. conf.researchr.org/track/icse-2026/icse-2026-artifact-evaluation",
 "Zenodo: size limitations and general policies. help.zenodo.org; about.zenodo.org/policies",
 "Heumüller, Nielebock, Krüger, Ortmeier: Publish or perish, but do not forget your software artifacts. EMSE 25 (2020)",
 "Hermann, Winter, Siegmund: Community Expectations for Research Artifacts and Evaluation Processes. ESEC/FSE (2020)",
 "Collberg, Proebsting: Repeatability in Computer Systems Research. CACM 59(3) (2016)",
 "Felter, Ferreira, Rajamony, Rubio: An Updated Performance Comparison of Virtual Machines and Linux Containers. IBM Research (2014)",
 "DFG: Guidelines for Safeguarding Good Research Practice, Code of Conduct (2019); 10-year archiving as implemented e.g. by Univ. Potsdam",
 "Anonymous GitHub (anonymous.4open.science). github.com/tdurieux/anonymous_github",
 "OOPSLA 2013 AEC nomination letter; Krishnamurthi: Artifact Evaluation for Software Conferences. cs.brown.edu/~sk/Memos/Conference-Artifact-Evaluation",
 "SOSP 2021 AEC call (sigops.org); APLAS 2022 AEC call",
 "SLE 2025 Artifact Evaluation. sleconf.org/2025/ArtifactEvaluation.html",
 "Miller: A Scientist's Nightmare: Software Problem Leads to Five Retractions. Science 314 (2006)",
 "Herndon, Ash, Pollin: Does High Public Debt Consistently Stifle Economic Growth? A Critique of Reinhart and Rogoff (2013)",
 "The Register: Python glob ordering bug in NMR scripts (2019-10-15)",
 "POPL 2023 Artifact Evaluation. popl23.sigplan.org/track/POPL-2023-artifact-evaluation",
 "ICSE Open Science Policy (ICSE 2023 technical track; ICSE 2027 Q&A). conf.researchr.org",
 "European Commission: Open science in Horizon Europe. rea.ec.europa.eu/open-science_en",
 "Artifact evaluation calls of CAV 2025, FORMATS 2022, QEST 2023",
 "MobiSys 2026 and IISWC 2026 artifact evaluation calls",
 "OSDI 2020 artifact evaluation report (SOSP 2019 Azure credits). sysartifacts.github.io/osdi2020/report",
 "Smith, Niemeyer, Katz et al.: Journal of Open Source Software (JOSS): design and first-year review. PeerJ CS (2018)",
 "NERC/EIDC data embargo FAQ (eidc.ac.uk/help/faq/embargos); telescope proprietary periods",
 "Princeton PICSciE: What's an RSE? (2018); Aalto Scientific Computing RSE",
]

SLIDE = {  # slide number -> [(old, new)] in visible text
 32: [("too big for Zenodo", "too big for Zenodo [4]")],
 38: [("e.g. a 30-minute setup", "e.g. a 30-minute setup [3]")],
}
NOTES = {
 17: [("(Heumüller: 48 % of ICSE papers linked runnable software)", "(Heumüller: 48 % of ICSE papers linked runnable software [5])")],
 19: [("are slower, which matters for performance claims.", "are slower, which matters for performance claims [8]."),
      ("but archived is not the same as still runnable.", "but archived is not the same as still runnable [9].")],
 20: [("each conference's AEC sets its own concrete rules, which is exactly why standards differ.",
       "each conference's AEC sets its own concrete rules, which is exactly why standards differ [1, 6].")],
 21: [("can archive GitHub releases automatically; ICSE now requires such an archive for the Available badge.",
       "can archive GitHub releases automatically [4]; ICSE now requires such an archive for the Available badge [3]."),
      ("but its links expire afterwards.", "but its links expire afterwards [10]."),
      ("ACM allows proxy data in place of proprietary data.", "ACM allows proxy data in place of proprietary data [1].")],
 22: [("(familiar with current tools, training the next generation).", "(familiar with current tools, training the next generation) [11]."),
      ("about 8-20 hours per artifact; ICSE's 30-minute rule only covers setup.", "about 8-20 hours per artifact [12]; ICSE's 30-minute rule only covers setup [3]."),
      ("where reviewers and authors can fix setup problems.", "where reviewers and authors can fix setup problems [13].")],
 24: [("treated as honest error.", "treated as honest error [14]."), ("found after the data was shared.", "found after the data was shared [15]."),
      ("some reviewers prefer availability over quality.", "some reviewers prefer availability over quality [6].")],
 25: [("no measured scooping rate.", "no measured scooping rate [24].")],
 26: [("to protect reviewers' machines.", "to protect reviewers' machines [17].")],
 27: [("archived is not runnable.", "archived is not runnable [9]."), ("short-cycle maintenance funding.", "short-cycle maintenance funding [25]."),
      ("source code most stable.", "source code most stable (week-1 discussion, not a source).")],
 28: [("wrong on another.", "wrong on another [16].")],
 30: [("(SOSP 2021, APLAS 2022).", "[12].")],
 32: [("up to 200 GB on request.", "up to 200 GB on request [4].")],
 33: [("only a statement of why not.", "only a statement of why not [18].")],
 34: [("Proxies allowed for proprietary data.", "Proxies allowed for proprietary data [1, 3].")],
 35: [("justified in a data management plan.", "justified in a data management plan [19].")],
 36: [("each AEC decides.", "each AEC decides [1]."), ("e.g. MATLAB Online.", "e.g. MATLAB Online [20].")],
 37: [("ECOOP/OOPSLA 2013;", "ECOOP/OOPSLA 2013 [11];"), ("(Hermann et al.).", "[6].")],
 38: [("makes acceptance unlikely;", "makes acceptance unlikely [3];"), ("reduced representative scope.", "reduced representative scope [21].")],
 39: [("proved hard to use;", "proved hard to use [22];"), ("provide remote access.", "provide remote access [21].")],
 40: [("'kick the tires' phase;", "'kick the tires' phase [13];"), ("answers within 72 h.", "answers within 72 h [3].")],
 42: [("- JOSS;", "- JOSS [23];")],
}

src, dst = sys.argv[1], sys.argv[2]
p = Presentation(src)
missing = []
for n, reps in SLIDE.items():
    for old, new in reps:
        hit = False
        for sh in p.slides[n - 1].shapes:
            if not sh.has_text_frame: continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if old in r.text: r.text = r.text.replace(old, new); hit = True
        if not hit: missing.append((n, old))
for n, reps in NOTES.items():
    tf = p.slides[n - 1].notes_slide.notes_text_frame
    t = tf.text
    for old, new in reps:
        if old in t: t = t.replace(old, new)
        else: missing.append((n, old))
    tf.text = t
# references slide
ref = next(s for s in p.slides if any(sh.has_text_frame and sh.text_frame.text.startswith("[1] ACM") for sh in s.shapes))
box = next(sh for sh in ref.shapes if sh.has_text_frame and sh.text_frame.text.startswith("[1] ACM"))
box.top = box.top - Pt(10); box.height = Pt(420)
tf = box.text_frame; tf.word_wrap = True
for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
first = tf.paragraphs[0]
for r in list(first.runs)[1:]: r._r.getparent().remove(r._r)
for i, txt in enumerate(REFS, 1):
    para = first if i == 1 else tf.add_paragraph()
    para.space_after = Pt(1)
    run = para.runs[0] if (i == 1 and para.runs) else para.add_run()
    run.text = f"[{i}] {txt}"; run.font.size = Pt(10.5); run.font.name = "Arial"; run.font.bold = False
p.save(dst)
print("missing:", missing)
