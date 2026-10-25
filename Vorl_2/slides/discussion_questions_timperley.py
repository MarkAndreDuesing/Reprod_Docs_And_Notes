"""Paper 2 (Timperley) discussion questions: single source for the notes file and the slides.

Usage: python3 discussion_questions_timperley.py <in.pptx> <out.pptx> <out.md>
Inserts the core question slides after the slide titled "Paper 2 · Discussion (Kevin)",
then the spares (hidden). The existing idea-dump slides follow and are hidden.
"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

Q = [  # (title, question, follow-ups, example answers, my notes)
 ("C1: Fear of mistakes being found",
  "Creators name the fear of mistakes in their analysis being discovered as a reason not to share. But isn't finding mistakes exactly the point of scientific review?",
  ["How much personal accountability for their code and analysis should researchers expect to carry once it is public?",
   "The paper lists this fear next to reputation damage from poor code quality. Does that weigh heavier on PhD students' prototype code?"],
  ["Not legitimate: errors found early are cheaper than errors found after others have built on them; hiding them only delays the damage.",
   "Understandable: incentives punish visible mistakes (retraction, reputation), while unshared mistakes cost the author nothing.",
   "Middle ground: an honest bug should be treated differently from a careless or misleading result, but today's system rarely makes that distinction."],
  ["Chang 2006: an inherited lab script flipped two data columns, leading to 5 retractions (3 in Science); treated as an honest error.",
   "Hermann et al. (this week): an AEC found hard-coded outputs and the paper was withdrawn.",
   "Reinhart & Rogoff 2013: a student's request for the spreadsheet revealed a range error that changed the headline result."]),
 ("C1: Holding back artifacts",
  "Some creators hold back artifacts that still have untapped publication value, e.g. to protect a PhD student from being scooped. Does this slow down research as a whole, and how far should researchers accept their work being built upon or undercut for scientific progress?",
  ["One respondent says whole research areas died because early work had no artifacts available.",
   "Others don't share because only a small group works on the topic. Is \"nobody else will use it\" a valid reason?"],
  ["Protect: a PhD has a fixed duration; being overtaken with your own tool is a real career risk.",
   "Share: the paper already establishes priority, and withholding hides the evidence behind its claims.",
   "Compromise: delaying release instead of withholding it (possible answer, not from the paper)."],
  ["Data-heavy fields use embargoes: telescope data 12-18 months, NERC up to 2 years.",
   "No measured rate of how often scooping actually happens; it is a perceived risk.",
   "Week 1: Collberg & Proebsting also listed competitors' advantage as a reason not to share."]),
 ("C3: Decay and maintenance",
  "Bugs, changing use cases and \"inevitable bit rot\" mean artifacts need ongoing maintenance, usually by students who later leave. Is software decay unavoidable as technology moves on? What practical ways are there to limit it?",
  ["Maintenance is an ongoing investment of time and resources, and respondents say only well-funded groups manage it. Who should carry that cost?",
   "Should an artifact only have to work at publication and review time, or for a set period?"],
  ["Unavoidable: dependencies, compilers and operating systems change; you can slow decay, not stop it.",
   "Limitable: archived VM or container images, pinned versions, few dependencies, handover when students leave.",
   "Cost: the group or university (it is research infrastructure) vs. nobody (archive it and mark it unmaintained)."],
  ["DFG: code and data behind a publication archived for at least 10 years; archived is not the same as runnable.",
   "Research software engineers keep code alive inside institutions (Princeton, Aalto).",
   "Maintenance funding (CZI open-source fund, DFG calls) runs in short cycles.",
   "Week 1: costly constant maintenance; Docker may not exist in 40 years; source code may be the most stable format."]),
 ("C4: One-click reproduction scripts",
  "Some reviewers suggest automated scripts to get around tacit knowledge, while others warn that a script is a black box that can reproduce the authors' mistakes. Is a one-click reproduction script good or bad for scientific scrutiny?",
  ["If a script does not help a reviewer understand why the results come out, what would?",
   "One reviewer found it \"easy to reproduce the experiments, but difficult to understand where should I add more case studies\"."],
  ["Good: anyone can re-run it without setup knowledge; without scripts many artifacts would not run at all.",
   "Bad: the reviewer checks that the numbers come out, not whether the method behind them is sound, which creates false confidence.",
   "What helps understanding: documented steps, small examples, a description of how each result is produced."],
  ["Python file-order case (2019): NMR scripts gave correct values on one OS and wrong ones on another; ~150 studies possibly affected.",
   "Push-button replication is an established first step, followed by an independent re-implementation."]),
 ("C9: Unshareable industry data",
  "One creator was criticised for not sharing data from an industry-track paper and called this a \"ritualistic blanket argument\" for judging papers. Should paper reviewers be allowed to criticise a paper for not sharing data that its industry partner won't release?",
  ["Should there be formal exceptions to sharing for such cases?",
   "One partner refused even anonymised data, fearing liability lawsuits. What could a paper offer instead?"],
  ["Allow criticism: readers cannot check the claims, so the evidence is weaker whatever the reason.",
   "Don't: it punishes valuable industry collaboration for something the authors cannot control.",
   "Middle: state the reason openly; offer proxy data or access for reviewers only."],
  ["ICSE: data sharing not required for acceptance, only a statement of why not.",
   "ACM 'Complete' criterion allows proxies for proprietary data.",
   "The Evaluated badges can be earned with reviewer-only access; Available needs public access."]),
 ("C9: Licensed software",
  "Users and reviewers report artifacts that need paid licences for auxiliary tools, or proprietary components such as Windows or a commercial IDE. Should research tools that depend on paid commercial software be accepted as artifacts at all?",
  ["Example: your artifact needs MATLAB, and the reviewer has no licence. Whose problem is that?",
   "Reviewers see \"no obvious pathway\" here. Is rejecting such artifacts fair, or should the reviewer or the conference arrange access?"],
  ["Accept: much real-world SE happens on commercial platforms; excluding them biases research toward open tools.",
   "Restrict: if others cannot run or check it, the artifact does not serve its purpose.",
   "It depends: is the commercial tool the subject of the research, or a replaceable helper?"],
  ["CAV, FORMATS and QEST handle it case by case: authors contact the chairs; FORMATS 2022 accepted MATLAB Online.",
   "Reviewers are not expected to buy licences.",
   "Paper: one user decompiled a tool to extend it and then could not share the extension because of licences."]),
 ("C10: Students as reviewers",
  "Artifact review is time-consuming and poorly rewarded, so \"we usually just make students do AEC\". Is that a good training opportunity, or does leaving it to comparatively inexperienced students undermine how rigorous the badges are and how seriously they are taken?",
  ["The paper warns this \"may lead to inconsistent and low-quality reviews\".",
   "What would a badge need to certify for you to trust it when choosing a tool for your own work?"],
  ["Meaningful: checking that an artifact runs and matches the paper does not need seniority; students are close to the code and typical future users.",
   "Weak: without experience and time, students may not judge the methodology or dare to reject; badges certify \"it ran once\"."],
  ["Students chosen deliberately from ECOOP/OOPSLA 2013: familiar with current tools, and trains the next generation.",
   "Time budget about 8-20 hours per artifact (SOSP 2021, APLAS 2022).",
   "Hermann et al.: most AEC members serve only once."]),
 ("C12: Artifacts that need author support",
  "Reviewers say a minor installation glitch shouldn't mean rejection, and that what matters is the authors' willingness to help resolve issues. If an artifact only works with help from its authors, does it deserve a badge, or should there be a separate category for artifacts that fail on their own but are actively supported?",
  ["Support can't last forever: authors become unresponsive, students leave.",
   "Should evaluation be a one-shot test or an iterative dialogue with the authors?"],
  ["Same badge: the review fixed the problems, so later users benefit too.",
   "No badge: a badge should mean \"works for a stranger\"; support disappears when the student graduates.",
   "Separate badge: honest and rewards responsive authors, but may encourage relying on support instead of fixing documentation."],
  ["Most venues have a 'kick the tires' phase and response windows (SLE: 5 days); ICSE 2026 requires answers within 72 h."]),
]
SPARES = [
 ("C2: Performance claims",
  "Reviewers note that in a VM \"it is not really possible to accurately assess performance\", and containers don't always stop software from degrading. How can artifact reviewers verify performance claims if they rarely have the authors' hardware?",
  ["Example: the paper claims a 3x speed-up, and the reviewer runs it in a VM on a laptop. What should count as reproduced?",
   "Should performance claims be left out of artifact evaluation?"],
  ["Check trends instead of exact numbers vs. leave performance out and certify only functionality."],
  ["SPLASH 2026 asks authors to say how reviewers on other hardware can recognise the high-level results.",
   "Containers run near native speed; VMs are slower (Felter et al. 2014)."]),
 ("C2: Running untrusted artifacts",
  "Running untrusted artifacts is a security risk for reviewers. Should reviewers be allowed to refuse artifacts that need root access or invasive libraries?",
  [], ["Yes, safety first vs. then system-level research becomes unevaluable."],
  ["POPL recommends VM images partly to protect reviewers' machines."]),
 ("C11: Artifacts too large or heavy to review",
  "Reviewers struggle when a VM image doesn't fit on their disk, and some experiments (cloud systems, fuzzers, program repair) need hardware they don't have. Are such artifacts simply unreviewable under the current process?",
  [], ["Weaker badge for what could be checked vs. a different kind of evaluation."],
  ["MobiSys: reviewers and authors agree on a reduced but representative scope; IISWC: authors provide remote access. Avoid 'who provides compute' (teammate's recommendations)."]),
 ("C7: Dead artifact links",
  "About 14 % of artifact links were already dead, more for older papers. Should conferences re-check artifact links a few years after publication and mark dead ones?",
  ["How would you organise and pay for this?"], ["Keeps badges honest vs. nobody has the mandate or budget."],
  ["Klein et al. 2014: 13-22 % of web references in articles suffer link rot."]),
 ("RQ1: Rewards",
  "Respondents want sharing \"encouraged and rewarded\". What reward would most motivate you as a researcher?",
  [], ["Career credit (hiring, citations) vs. direct benefit (users, collaborations)."],
  ["JOSS: citable software papers. Possible overlap with the teammate's recommendations."]),
]

def markdown():
    out = ["# Timperley et al. (2021): Discussion Questions (Challenges, §4.2)", "",
           "Generated from `Vorl_2/slides/discussion_questions_timperley.py` (single source for slides and notes).", "",
           "Core questions come from my selected (green) questions, lightly reworded and anchored in the paper. "
           "Follow-ups go on the slide as sub-points. Example answers and my notes go in the speaker notes only; "
           "the notes draw on `timperley-research.md` and are not part of the questions.", ""]
    def block(items, hdr):
        out.append(f"## {hdr}"); out.append("")
        for i, (t, q, f, a, n) in enumerate(items, 1):
            out.extend([f"### {i}. {t}", "", f"**Question:** {q}", ""])
            if f: out.extend(["**Follow-ups:**"] + [f"- {x}" for x in f] + [""])
            out.extend(["**Example answers:**"] + [f"- {x}" for x in a] + [""])
            out.extend(["**My notes:**"] + [f"- {x}" for x in n] + [""])
    block(Q, "Core 8"); block(SPARES, "Spares (hidden slides)")
    return "\n".join(out)

def build(src, dst):
    G = RGBColor(0x00, 0x88, 0x3A); K = RGBColor(0x1A, 0x1A, 0x1A); M = RGBColor(0x5F, 0x63, 0x68)
    p = Presentation(src)
    slides = list(p.slides)
    def ttl(s): return next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
    anchor = next(s for s in slides if ttl(s).startswith("Paper 2 · Discussion (Kevin)"))
    dumps = [s for s in slides[slides.index(anchor) + 1:slides.index(anchor) + 4] if ttl(s) == "Paper 2 · Discussion"]
    ref = next(s for s in slides if ttl(s) == "General Summary" and slides.index(s) > 10)
    title_el = next(sh for sh in ref.shapes if sh.has_text_frame and sh.text_frame.text == "General Summary")._element
    num_el = next(sh for sh in ref.shapes if sh.is_placeholder)._element
    def add(item, hidden):
        t, q, f, a, n = item
        s = p.slides.add_slide(ref.slide_layout)
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        s.shapes._spTree.append(copy.deepcopy(title_el)); s.shapes[-1].text_frame.paragraphs[0].runs[0].text = "Paper 2 · Discussion"
        s.shapes._spTree.append(copy.deepcopy(num_el))
        def para(tf, first, text, size, bold=False, color=K, space=8, prefix=None):
            pa = tf.paragraphs[0] if first else tf.add_paragraph(); pa.space_after = Pt(space)
            if prefix:
                r = pa.add_run(); r.text = prefix; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = G; r.font.name = "Arial"
            r = pa.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Arial"
        box = s.shapes.add_textbox(Inches(.7), Inches(1.45), Inches(11.93), Inches(1.0)); tf = box.text_frame; tf.word_wrap = True
        para(tf, True, t, 28, True, G, 4)
        box = s.shapes.add_textbox(Inches(.7), Inches(2.25), Inches(11.93), Inches(2.4)); tf = box.text_frame; tf.word_wrap = True
        para(tf, True, q, 24 if len(q) < 260 else 22, True, K, 4)
        if f:
            h = 0.55 + 0.65 * len(f)
            pn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.7), Inches(7.0 - h), Inches(11.93), Inches(h - .2))
            pn.adjustments[0] = 0.08; pn.fill.solid(); pn.fill.fore_color.rgb = RGBColor(0xF1, 0xF3, 0xF2); pn.line.fill.background(); pn.shadow.inherit = False
            box = s.shapes.add_textbox(Inches(.9), Inches(7.0 - h + .1), Inches(11.5), Inches(h - .4)); tf = box.text_frame; tf.word_wrap = True
            for i, x in enumerate(f): para(tf, i == 0, x, 19, False, K, 6, "↳ ")
        s.notes_slide.notes_text_frame.text = ("EXAMPLE ANSWERS\n" + "\n".join("- " + x for x in a) +
                                               "\n\nMY NOTES (background, not on slide)\n" + "\n".join("- " + x for x in n))
        if hidden: s._element.set("show", "0")
        return s
    new = [add(x, False) for x in Q] + [add(x, True) for x in SPARES]
    for d in dumps: d._element.set("show", "0")
    lst = p.slides._sldIdLst
    def eid(sl): return next(e for e in lst if p.part.rels[e.rId]._target is sl.part)
    prev = eid(anchor)
    for s in new:
        e = eid(s); lst.remove(e); prev.addnext(e); prev = e
    p.save(dst)
    print("slides", len(p.slides), "dumps hidden", len(dumps))

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
    open(sys.argv[3], "w").write(markdown() + "\n")
