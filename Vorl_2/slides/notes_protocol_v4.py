"""Adds speaking protocols to slides 15 and 17-22 of the team deck (v4). Existing notes are kept below."""
import sys
from pptx import Presentation
P = {
15: """SPEAKING PROTOCOL (~30 s)
- "Two research questions. RQ1 is the baseline: does the SE community see value in artifacts at all? Only if yes is it worth asking RQ2: what gets in the way?"
- "RQ2 is split by perspective: creators who make artifacts, users who try to run or build on someone else's, and reviewers on artifact evaluation committees."
- Point out: the same person can be all three; most respondents were creators and users, only about a fifth had reviewed artifacts.
- Hand-over: "Let's start with RQ1." """,
17: """SPEAKING PROTOCOL (~45 s)
- "Short answer to RQ1: yes. Respondents see artifacts as valuable, say not every paper needs one, but that sharing should be encouraged and rewarded."
- "Their main argument: without the artifact, replicating a study is hard and unreliable, and you cannot meaningfully evaluate the approach."
- Backed by the publication survey: almost two thirds of papers at ICSE, FSE, ASE and EMSE (2014-2018) link an artifact, rising from about half to about 70 % over the five years.
- Walk through the 4 tiles briefly: supporting evidence (readers trust the claims more), community norm ("it's good practice"), reuse (tools/data others extend), understanding (details that don't fit in a paper).
- Worth saying out loud: "artifact" here means any linked material, so this counts companion websites and survey sheets too, which is why the share is higher than last week's numbers (Heumüller: 48 % of ICSE papers linked runnable software).
- Bridge: "So people value artifacts. Why are they still often missing or unusable? That's RQ2." """,
18: """SPEAKING PROTOCOL (~40 s)
- "These are the 12 challenges they found. They come from a thematic analysis of the open survey answers: so this shows which problems exist, not how often each occurs."
- Explain the symbols: square = creators mention it, triangle = users, star = reviewers.
- Point out the pattern: portability, tacit knowledge and external constraints hit all three groups; the last three (incentives, technical obstacles, communication) are reviewer-only; 'not worth it' and double-blind review are creator-only.
- Optional framing from the paper: the problems boil down to communication, reward/social system, time ('bitrot') and the technical side of the artifacts.
- "I'll go through them in groups of three, focusing on what's new or surprising." """,
19: """SPEAKING PROTOCOL (~60 s)
C1 Not worth it
- "Biggest creator issue: preparing an artifact costs time that could go into the next paper, and it does little for your career."
- List the risks quickly: bad code hurts reputation, being scooped, mistakes being found, maintenance burden, nobody using it.
- Mention the counterpoint from a respondent: whole research areas have died because early work had no artifacts and re-implementing it was too costly.
C2 Portability
- "It works on my machine": creators can't foresee every environment; users find hard-wired paths and dependencies that no longer exist (building and missing dependencies are among the most common user problems).
- Reviewers: installing is the biggest hurdle. VMs protect the reviewer's machine but are huge and distort performance; containers help, but don't stop software from decaying over time.
C3 Maintenance
- "Artifacts need ongoing care: bugs, changing use cases, bit rot. Usually a PhD student maintains it, and then graduates."
- Users see the result: old library versions that are harder and harder to get.
BACKGROUND worth saying if time
- Container vs VM: containers run almost at native speed; VMs are slower, which matters for performance claims.
- Germany's DFG expects code and data behind a publication to be archived for 10 years, but archived is not the same as still runnable.""",
20: """SPEAKING PROTOCOL (~60 s)
C4 Tacit knowledge
- "Know-how that only lives in the developer's head: which order to run the scripts, which paths to change, which parameters matter."
- Users: missing usage and build instructions are the most and third-most common problems.
- Reviewers: some had to read the paper to fill gaps in the instructions. One-click scripts help, but act as a black box and may reproduce the authors' mistakes.
C5 Doesn't fit purpose
- Users only: the artifact doesn't do what the paper claims or only works for very narrow cases, so people end up modifying or re-implementing it.
C6 No standards / guidelines
- Creators don't know what to include or how to package it. Reviewers find the criteria and ACM badge descriptions vague.
- Key open question: should reviewers just run the scripts or dig deep into the code? And should a broken artifact affect the paper's acceptance, or is artifact review just a "rubber stamp"?
BACKGROUND worth saying
- ACM only defines the badges (Functional, Reusable, Available, Reproduced, Replicated); each conference's AEC sets its own concrete rules, which is exactly why standards differ.""",
21: """SPEAKING PROTOCOL (~60 s)
C7 Hosting
- "Where do you put an artifact so it's still there in ten years, can hold large data, and is free? GitHub, Dropbox or Google Drive are not archives."
- Users: dead links are the second most common problem. In the publication survey about 14 % of links were already dead in 2019, about a quarter for 2014 papers.
C8 Double-blind review
- Sharing with the submission means anonymising the artifact and finding anonymous hosting: extra hassle, so some skip sharing.
C9 External constraints
- Things authors can't control: IP and contracts with industry partners, fear of liability, privacy (even anonymised interview data can be re-identified).
- One author was criticised for not sharing industry data and called that a "ritualistic blanket argument".
- Users: licences block extending tools; some need paid licences. Reviewers: no licence for e.g. Windows or a commercial IDE, and no clear rule what to do.
BACKGROUND worth saying
- Zenodo (run by CERN) is free, gives a DOI, holds up to 50 GB per record, and can archive GitHub releases automatically; ICSE now requires such an archive for the Available badge.
- Anonymous GitHub mirrors a repository without names for review, but its links expire afterwards.
- ACM allows proxy data in place of proprietary data.""",
22: """SPEAKING PROTOCOL (~60 s)
Start with the context line: "First some numbers: at FSE 2015-2018, only about a third of papers with artifacts were submitted to artifact evaluation, and 86 % of those passed. So participation is low, and passing is easy."
C10 No reviewer incentives
- Reviewing artifacts takes a lot of time and counts less than paper reviewing. "We usually just make students do AEC": good training, but it signals artifacts are secondary, and may lead to inconsistent reviews.
C11 Technical obstacles
- Long runtimes and little compute (reviewers have roughly a day per artifact), huge VM images that don't fit on a laptop, and some experiments (cloud systems, fuzzing, program repair) are practically impossible to re-run.
C12 Limited communication
- Some AECs review in a single round, so a typo in the install instructions can mean rejection. Reviewers value authors who help; unresponsive authors leave them waiting for days.
Wrap-up line: "Reviewers face the same problems as users (portability, docs) and as creators (no guidelines, no reward). That's the end of the results; next, the discussion."
BACKGROUND worth saying
- Artifact committees were deliberately staffed with senior PhD students from 2013 on (familiar with current tools, training the next generation).
- Typical review budget is about 8-20 hours per artifact; ICSE's 30-minute rule only covers setup.
- Most venues have a "kick the tires" phase where reviewers and authors can fix setup problems.""",
}
src, dst = sys.argv[1], sys.argv[2]
p = Presentation(src)
for n, text in P.items():
    tf = p.slides[n - 1].notes_slide.notes_text_frame
    old = tf.text.strip()
    tf.text = text + ("\n\nEARLIER NOTES\n" + old if old else "")
p.save(dst)
print("done")
