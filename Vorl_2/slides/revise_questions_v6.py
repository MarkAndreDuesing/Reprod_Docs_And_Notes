"""Third question revision on the team deck (v6): edits question slides in place, keeps the user's
tier order (Good / Further / hidden Backup), moves 'Hardware reviewers don't have' to Backup."""
import sys
from pptx import Presentation
from pptx.util import Pt

def notes(answers, background):
    return ("EXAMPLE ANSWERS\n" + "\n".join("- " + a for a in answers) +
            "\n\nBACKGROUND (sources named inline; not on slide)\n" + "\n".join("- " + b for b in background))

# subheader -> new content (None = keep that field)
EDIT = {
 "C4 Tacit knowledge: One-click reproduction scripts": dict(
   notes=notes(["Good: anyone can re-run it without setup knowledge.",
                "Bad: checks that the numbers come out, not that the method is sound."],
               ["Follow-up (from the dropped tacit-knowledge slide): Could someone reproduce your bachelor thesis results from your repository alone?",
                "NMR script case (The Register, 2019): Python's glob returns files in an OS-dependent order; widely used chemistry scripts were correct on Windows/older macOS but wrong on newer macOS and Ubuntu, possibly affecting ~150 studies. A script that 'works' can still be wrong.",
                "'Push-button replication' (3ie, development economics) is used only as a first check, followed by an independent re-implementation."])),
 "C4 Tacit knowledge: Reviewer effort": dict(
   question="How much time, compute and money should a reviewer invest before giving up on an artifact?",
   example="reproducing it needs a paid cloud GPU for a day, or fixing a broken dependency yourself",
   notes=notes(["Little: it should work for a stranger; reviewers are unpaid volunteers.",
                "More: reviewers are often the only ones who ever try; one fix helps every later user."],
               ["Typical budget: about 8-20 hours per artifact (SOSP 2021 AEC call: 10-20 h, 2-4 artifacts per member; APLAS 2022: ~8 h, 3-4 per member).",
                "No known dedicated budget for artifact evaluation; where cloud costs arise, venues shift them to authors (SOSP 2023 allows public cloud only if authors pay; the Programming journal suggests authors pay for a cloud instance reviewers can use)."])),
 "C9 External constraints: Licensed software": dict(
   notes=notes(["Accept: much real-world SE runs on commercial platforms.",
                "Restrict: if others cannot run it, it misses its purpose.",
                "Depends whether the commercial tool is the subject of the research or a replaceable helper."],
               ["Link to the previous slide: 'We just asked how open research software should be; here the limit is not the authors' choice but a licence.'",
                "Who pays: artifact evaluation committees are unpaid volunteers set up by the organisers; no venue states a budget for AE or for reviewer licences (no source found; likely, not verified).",
                "Practice: CAV asks authors who cannot ship licensed software such as MATLAB in a VM/Docker image to contact the chairs for a case-by-case solution; FORMATS 2022 accepted a MATLAB Drive folder runnable in MATLAB Online.",
                "ACM's badge policy says nothing about reviewer licences; each committee decides.",
                "Possible follow-up (from the dropped 'formal exceptions' slide): reviewers get temporary access under NDA, but readers never see the artifact: would that badge convince you?"])),
 "C2 Portability: Running untrusted artifacts": dict(
   context="Some artifacts need root access or invasive system libraries; reviewers use VMs to protect their own machines.",
   notes=notes(["Yes: safety first; authors must package safely (VM, container).",
                "No: system-level research (OS, drivers, security tools) would become unevaluable."],
               ["From the paper: users report artifacts that 'require the user to be root' (p. 18); reviewers cite the 'requirement to install complex and invasive program/system libraries' (p. 21); VMs help 'avoid the security risks associated with executing untrusted code on the reviewer's host machine' (p. 21).",
                "POPL's artifact call recommends VM images partly because they 'protect reviewers from malfunctioning artifacts damaging their computer'."])),
 "C3 Maintenance: Decay and maintenance": dict(
   example="a tool that only runs on Python 2, which has received no updates since 2020",
   you="Should your thesis tool still run when the next student picks it up three years later?",
   notes=notes(["Unavoidable: dependencies, compilers and operating systems change; you can slow decay, not stop it.",
                "Limitable: archived images, pinned versions, few dependencies, handover when students leave.",
                "Cost: group or university vs. nobody (archive and mark unmaintained)."],
               ["Python 2 reached end of life on 1 Jan 2020 and Linux distributions such as Arch removed it; the metabarcoding tool OBITools (Python 2) is no longer maintainable and is kept runnable only via a Docker image its maintainers provide.",
                "Docker Hub announced in 2020 that images of free accounts not pulled for 6 months would be deleted; enforcement was paused after protests. Even 'frozen' images depend on a company's policy.",
                "NMR scripts (2019): correct for years, then wrong on newer operating systems.",
                "Germany's DFG expects code and data behind a publication to be archived for 10 years; archived is not the same as runnable.",
                "Research software engineers (e.g. Princeton, Aalto) are permanent staff who keep research code alive when students leave."])),
 "C7 Hosting: Very large artifacts": dict(
   title="C7 Hosting / C11 Technical obstacles: Unreviewable artifacts",
   context="Some artifacts are too large or need hardware no reviewer has.",
   question="If no reviewer can run an artifact, should it still be eligible for a badge, and what would the badge mean?",
   example="several terabytes of data, or an experiment that needs eight GPUs for a week",
   you="Would you trust a badge awarded after only a scaled-down run?",
   notes=notes(["Yes, with a badge that says what was actually checked.",
                "No: a check nobody could run certifies nothing.",
                "Alternative: a reduced but representative run, or authors provide access to their hardware."],
               ["From the paper: large VM images that don't fit on a reviewer's disk; cloud systems, fuzzers and program repair need hardware reviewers don't have (pp. 23-24).",
                "Current practice: authors give reviewers anonymous remote (SSH) access to special hardware (SenSys 2024, SACMAT 2026, IEEE S&P 2026, MICRO 2026). MICRO 2026 may decline to evaluate artifacts if no reviewer can access the hardware, i.e. they become unreviewable.",
                "MobiSys 2026: reviewers and authors may agree on a reduced but representative scope.",
                "Zenodo holds 50 GB per record (200 GB on request), so multi-terabyte data needs other repositories.",
                "Hand-over: the paper recommends that ACM/IEEE fund compute for reviewers (Rec 6/7, Minh Thi's section), while venues today shift the burden to authors."])),
 "C9 External constraints: Unshareable industry data": dict(
   question="Should top venues accept industry studies whose data no one else can ever check, or would refusing them push industry research out?",
   you="If your thesis were done with a company and its data couldn't be shared, would you still want to publish at a top venue?",
   notes=notes(["Accept: industry data shows how software is really built; excluding it narrows SE research.",
                "Refuse: results nobody can check are weaker evidence, whatever the reason."],
               ["From the paper: partners refused even anonymised data, fearing liability lawsuits (p. 17); one author was criticised for not sharing industry data and called it a 'ritualistic blanket argument' (pp. 17-18).",
                "ICSE's open science policy: data sharing is not required for acceptance; authors state why they cannot share in a data availability section.",
                "ACM allows proxy data in place of proprietary data; the Evaluated badges can be earned with reviewer-only access."])),
 "C11 Technical obstacles: Fixed time limits": dict(
   question="Is one flat time limit fair across very different artifacts, or should it depend on the artifact type?",
   example="a small library and a three-day fuzzing experiment under the same 30-minute setup rule [1]",
   notes=notes(["Flat: simple and predictable; future users won't invest more time either.",
                "By type: heavy artifacts (fuzzing, large data) are penalised by a flat limit."],
               ["[1] ICSE 2026 artifact evaluation call: if installation and configuration take more than 30 minutes, acceptance is unlikely.",
                "The 30 minutes is a setup guideline at many venues (OSDI 'e.g., within 30 minutes', OOPSLA 'roughly 30 minutes'); the whole review takes about 8-20 h per artifact; NDSS 2027 asks that full experiments take at most 1 day on a desktop.",
                "No venue found that varies limits by artifact type; heavy artifacts instead get a negotiated reduced scope (MobiSys) or remote access."])),
 "C12 Limited communication: One-shot or iterative review": dict(
   title="C12 Limited communication: Purpose of artifact review",
   context="Reviewers value authors who help; the other paper finds reviewers want reproducible results, users want reusable tools.",
   question="Should artifact review confirm that the results could be reproduced once, or that the artifact works as a foundation others can build on?",
   example="the artifact only runs after three emails with the authors: badge or not?",
   you="If you started a thesis on a badged tool, what would you expect the badge to guarantee?",
   notes=notes(["Reproduced once: that's what peer review of claims needs; reuse is the authors' business.",
                "Foundation: a badge readers rely on should mean others can build on it without the authors."],
               ["Hermann et al. (2020, Paper 1, p. 479): 'Reviewers want replicability, users want reusability' (old ACM terms: their 'replicability' is today's reproducibility). Also: 34 users saw no difference between evaluated and unevaluated artifacts, 20 did (p. 477).",
                "Follow-up: should evaluation be a one-shot test or a dialogue with the authors? Later users won't have the authors' help.",
                "Follow-up (from the dropped support slide): Would you answer emails about your thesis code in three years?",
                "Practice: most venues have a 'kick the tires' phase with author responses (SLE 2025: 5 days; ICSE 2026: 72 h)."])),
 "RQ1: Rewards": dict(
   question="The paper names four reasons people share artifacts; which reward would strengthen them enough to make the effort worth it?",
   example="citations of the artifact itself, a badge on the paper, artifacts counted in hiring, an award",
   notes=notes(["Career credit: counted in hiring and promotion.",
                "Visibility: badges and citations of the artifact.",
                "Norm: no reward needed if sharing is simply expected."],
               ["The four RQ1 motivations: supporting evidence, community norm, facilitating reuse, improving understanding (Table 3).",
                "JOSS (Journal of Open Source Software) publishes short, citable papers about research software so software work counts for careers.",
                "Hand-over: 'The paper itself makes recommendations on exactly this; Minh Thi will present them next.'"])),
}

def set_slide(s, c):
    shapes = [sh for sh in s.shapes if sh.has_text_frame]
    def find(pred): return next(sh for sh in shapes if pred(sh.text_frame.text))
    sub = find(lambda t: t.startswith(("C1", "C2", "C3", "C4", "C7", "C9", "C10", "C11", "C12", "RQ1")))
    order = [sh for sh in shapes if sh.text_frame.text.strip() and sh.text_frame.text not in ("‹#›",)]
    i = order.index(sub); ctx, q = order[i + 1], order[i + 2]
    box = find(lambda t: t.startswith("↳"))
    def put(sh, txt):
        runs = sh.text_frame.paragraphs[0].runs
        runs[0].text = txt
        for r in runs[1:]: r.text = ""
    if c.get("title"): put(sub, c["title"])
    if c.get("context"): put(ctx, c["context"])
    if c.get("question"): put(q, c["question"])
    for pa in box.text_frame.paragraphs:
        lab = pa.runs[0].text
        if "Example" in lab and c.get("example"): pa.runs[1].text = c["example"]
        if "You" in lab and c.get("you"): pa.runs[1].text = c["you"]
    if c.get("notes"): s.notes_slide.notes_text_frame.text = c["notes"]

src, dst = sys.argv[1], sys.argv[2]
p = Presentation(src)
def sub_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(tuple(k.split(":")[0].split(" ")[0] for k in EDIT)) and ":" in sh.text_frame.text:
            return sh.text_frame.text
done = set()
for s in p.slides:
    if s._element.get("show") == "0": continue
    t = sub_of(s)
    if t in EDIT: set_slide(s, EDIT[t]); done.add(t)
print("not found:", set(EDIT) - done)
# move hardware slide to backup (after the last hidden C12 backup slide) and hide it
lst = p.slides._sldIdLst
eid = lambda sl: next(e for e in lst if p.part.rels[e.rId]._target is sl.part)
hw = next(s for s in p.slides if sub_of(s) == "C11 Technical obstacles: Hardware reviewers don't have")
anchor = next(s for s in p.slides if sub_of(s) == "C12 Limited communication: Support after publication")
e = eid(hw); lst.remove(e); eid(anchor).addnext(e); hw._element.set("show", "0")
# references: only the ICSE call is cited on visible slides now
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("[1] ACM"):
            tf = sh.text_frame
            for pa in list(tf.paragraphs)[1:]: pa._p.getparent().remove(pa._p)
            runs = tf.paragraphs[0].runs
            runs[0].text = "[1] ICSE 2026 Artifact Evaluation, call for submissions. conf.researchr.org/track/icse-2026/icse-2026-artifact-evaluation"
            for r in runs[1:]: r.text = ""
p.save(dst)
